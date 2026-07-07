from flask import Flask, request, jsonify, render_template, send_file
from ultralytics import YOLO
import cv2
import numpy as np
import os
import sqlite3
from datetime import datetime
import json
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import time

app = Flask(__name__)

model = YOLO('yolov8n.pt')

os.makedirs('static', exist_ok=True)
os.makedirs('exports', exist_ok=True)

def init_db():
    conn = sqlite3.connect('history.db')
    cursor = conn.cursor()
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS requests (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            timestamp TEXT,
            count INTEGER,
            image_path TEXT,
            objects_detected TEXT,
            processing_time REAL,
            image_size TEXT,
            confidence_avg REAL
        )
    ''')
    conn.commit()
    conn.close()

init_db()

@app.route('/')
def index():
    return render_template('index.html')

@app.route('/process', methods=['POST'])
def process_image():
    file = request.files['image']
    start_time = time.time()
    
    img_bytes = np.frombuffer(file.read(), np.uint8)
    img = cv2.imdecode(img_bytes, cv2.IMREAD_COLOR)
    
    img_height, img_width = img.shape[:2]
    image_size = f"{img_width}x{img_height}"
    
    results = model(img)
    
    elephant_count = 0
    objects_info = []
    confidences = []
    
    for box in results[0].boxes:
        cls_id = int(box.cls[0].item())
        cls_name = model.names[cls_id]
        confidence = float(box.conf[0].item())
        
        if cls_name == 'elephant':
            elephant_count += 1
            x1, y1, x2, y2 = map(int, box.xyxy[0])
            confidences.append(confidence)
            
            # Рисуем рамку
            cv2.rectangle(img, (x1, y1), (x2, y2), (0, 255, 0), 3)
            label = f'Elephant {confidence:.2f}'
            cv2.putText(img, label, (x1, y1 - 10), 
                       cv2.FONT_HERSHEY_SIMPLEX, 0.8, (0, 255, 0), 2)
            
            objects_info.append({
                'class': cls_name,
                'confidence': round(confidence, 4),
                'bbox': [x1, y1, x2, y2]
            })
    
    processing_time = round(time.time() - start_time, 3)
    avg_confidence = round(np.mean(confidences), 4) if confidences else 0.0
    
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    result_filename = f'result_{timestamp}.jpg'
    result_path = os.path.join('static', result_filename)
    cv2.imwrite(result_path, img)
    
    conn = sqlite3.connect('history.db')
    cursor = conn.cursor()
    cursor.execute('''
        INSERT INTO requests (timestamp, count, image_path, objects_detected, 
                            processing_time, image_size, confidence_avg)
        VALUES (?, ?, ?, ?, ?, ?, ?)
    ''', (timestamp, elephant_count, result_path, json.dumps(objects_info),
          processing_time, image_size, avg_confidence))
    conn.commit()
    conn.close()
    
    return jsonify({
        'count': elephant_count,
        'image_url': result_path,
        'objects': objects_info,
        'processing_time': processing_time,
        'image_size': image_size,
        'avg_confidence': avg_confidence
    })

@app.route('/history')
def get_history():
    conn = sqlite3.connect('history.db')
    cursor = conn.cursor()
    cursor.execute('''
        SELECT id, timestamp, count, image_path, objects_detected, 
               processing_time, image_size, confidence_avg 
        FROM requests ORDER BY id DESC
    ''')
    rows = cursor.fetchall()
    conn.close()
    
    history = []
    for row in rows:
        history.append({
            'id': row[0],
            'timestamp': row[1],
            'count': row[2],
            'image_path': row[3],
            'objects': json.loads(row[4]) if row[4] else [],
            'processing_time': row[5],
            'image_size': row[6],
            'avg_confidence': row[7]
        })
    
    return jsonify(history)

@app.route('/export/json')
def export_json():
    conn = sqlite3.connect('history.db')
    cursor = conn.cursor()
    cursor.execute('''
        SELECT id, timestamp, count, image_path, objects_detected,
               processing_time, image_size, confidence_avg
        FROM requests ORDER BY id DESC
    ''')
    rows = cursor.fetchall()
    conn.close()
    
    data = []
    for row in rows:
        data.append({
            'id': row[0],
            'timestamp': row[1],
            'elephants_count': row[2],
            'image_path': row[3],
            'detected_objects': json.loads(row[4]) if row[4] else [],
            'processing_time_sec': row[5],
            'image_resolution': row[6],
            'avg_confidence': row[7]
        })
    
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    json_filename = f'exports/elephant_report_{timestamp}.json'
    
    with open(json_filename, 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    
    return send_file(json_filename, as_attachment=True, 
                    download_name=f'elephant_report_{timestamp}.json')

@app.route('/export/excel')
def export_excel():
    conn = sqlite3.connect('history.db')
    cursor = conn.cursor()
    cursor.execute('''
        SELECT id, timestamp, count, image_path, processing_time, 
               image_size, confidence_avg
        FROM requests ORDER BY id DESC
    ''')
    rows = cursor.fetchall()
    conn.close()
    
    wb = Workbook()
    
    ws1 = wb.active
    ws1.title = "Учет слонов"
    
    headers1 = ['ID', 'Дата и время', 'Количество слонов', 
                'Путь к изображению', 'Время обработки (сек)', 
                'Размер изображения', 'Средняя уверенность', 'Статус']
    ws1.append(headers1)
    
    for row in rows:
        status = 'Обнаружено' if row[2] > 0 else 'Не обнаружено'
        ws1.append([row[0], row[1], row[2], row[3], 
                   f"{row[4]:.3f}", row[5], f"{row[6]:.4f}", status])
    
    header_fill = PatternFill(start_color="2E5C8A", end_color="2E5C8A", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=12)
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'),
                        top=Side(style='thin'), bottom=Side(style='thin'))
    
    for cell in ws1[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.border = thin_border
        cell.alignment = Alignment(horizontal='center')
    
    for row in ws1.iter_rows(min_row=2, max_row=ws1.max_row, max_col=8):
        for cell in row:
            cell.border = thin_border
            if cell.column == 5:  # Время обработки
                cell.alignment = Alignment(horizontal='right')
                if cell.value:
                    try:
                        val = float(cell.value)
                        if val < 0.5:
                            cell.fill = PatternFill(start_color="C6EFCE", 
                                                   end_color="C6EFCE", fill_type="solid")
                        elif val < 1.0:
                            cell.fill = PatternFill(start_color="FFEB9C", 
                                                   end_color="FFEB9C", fill_type="solid")
                        else:
                            cell.fill = PatternFill(start_color="FFC7CE", 
                                                   end_color="FFC7CE", fill_type="solid")
                    except:
                        pass
    
    for column in ws1.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        ws1.column_dimensions[column_letter].width = adjusted_width
    
    ws2 = wb.create_sheet(title="Тесты производительности")
    
    headers2 = ['ID', 'Время обработки (сек)', 'Размер изображения', 
                'Количество объектов', 'Уверенность (%)', 'Производительность',
                'Объектов в секунду']
    ws2.append(headers2)
    
    for row in rows:
        objects_per_sec = round(row[2] / row[4], 2) if row[4] > 0 else 0
        performance = "Отлично" if row[4] < 0.5 else "Хорошо" if row[4] < 1.0 else "Нормально"
        ws2.append([row[0], f"{row[4]:.3f}", row[5], row[2], 
                   f"{row[6]*100:.2f}%", performance, objects_per_sec])
    
    for cell in ws2[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.border = thin_border
        cell.alignment = Alignment(horizontal='center')
    
    for row in ws2.iter_rows(min_row=2, max_row=ws2.max_row, max_col=7):
        for cell in row:
            cell.border = thin_border
    
    if rows:
        ws2.append([])
        ws2.append([])
        ws2.append(['СТАТИСТИКА ПРОИЗВОДИТЕЛЬНОСТИ'])
        
        total_images = len(rows)
        total_time = sum(row[4] for row in rows)
        avg_time = total_time / total_images if total_images > 0 else 0
        min_time = min(row[4] for row in rows)
        max_time = max(row[4] for row in rows)
        total_objects = sum(row[2] for row in rows)
        
        stats = [
            ['Всего обработано изображений:', total_images],
            ['Общее время обработки (сек):', f"{total_time:.3f}"],
            ['Среднее время обработки (сек):', f"{avg_time:.3f}"],
            ['Минимальное время (сек):', f"{min_time:.3f}"],
            ['Максимальное время (сек):', f"{max_time:.3f}"],
            ['Всего обнаружено слонов:', total_objects],
            ['Средняя скорость (изобр/сек):', f"{total_images/total_time:.2f}" if total_time > 0 else "N/A"]
        ]
        
        for stat in stats:
            ws2.append(stat)
    
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    excel_filename = f'exports/elephant_report_{timestamp}.xlsx'
    wb.save(excel_filename)
    
    return send_file(excel_filename, as_attachment=True,
                    download_name=f'elephant_report_{timestamp}.xlsx')

@app.route('/export/csv')
def export_csv():
    conn = sqlite3.connect('history.db')
    cursor = conn.cursor()
    cursor.execute('''
        SELECT id, timestamp, count, image_path, processing_time, 
               image_size, confidence_avg
        FROM requests ORDER BY id DESC
    ''')
    rows = cursor.fetchall()
    conn.close()
    
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    csv_filename = f'exports/elephant_report_{timestamp}.csv'
    
    with open(csv_filename, 'w', encoding='utf-8') as f:
        f.write('ID;Дата и время;Количество слонов;Путь к изображению;')
        f.write('Время обработки (сек);Размер изображения;Средняя уверенность\n')
        for row in rows:
            f.write(f'{row[0]};{row[1]};{row[2]};{row[3]};')
            f.write(f'{row[4]:.3f};{row[5]};{row[6]:.4f}\n')
    
    return send_file(csv_filename, as_attachment=True,
                    download_name=f'elephant_report_{timestamp}.csv')

@app.route('/stats')
def get_stats():
    conn = sqlite3.connect('history.db')
    cursor = conn.cursor()
    cursor.execute('''
        SELECT COUNT(*), SUM(count), AVG(processing_time), 
               MIN(processing_time), MAX(processing_time),
               SUM(processing_time)
        FROM requests
    ''')
    stats = cursor.fetchone()
    conn.close()
    
    return jsonify({
        'total_images': stats[0] or 0,
        'total_elephants': stats[1] or 0,
        'avg_processing_time': round(stats[2], 3) if stats[2] else 0,
        'min_processing_time': round(stats[3], 3) if stats[3] else 0,
        'max_processing_time': round(stats[4], 3) if stats[4] else 0,
        'total_time': round(stats[5], 3) if stats[5] else 0
    })

if __name__ == '__main__':
    app.run(debug=True, host='0.0.0.0', port=5000)